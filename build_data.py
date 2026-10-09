#!/usr/bin/env python3
"""Собирает data.js из таблицы xlsx (названия + гиперссылки на файлы в GitHub).
Запуск:  python3 tools/build_data.py путь/к/таблице.xlsx   (нужен openpyxl)
"""
import json, re, sys, datetime
from collections import OrderedDict
from urllib.parse import unquote, urlparse
from openpyxl import load_workbook

SRC = sys.argv[1] if len(sys.argv) > 1 else "table.xlsx"
OUT = sys.argv[2] if len(sys.argv) > 2 else "data.js"
ORDER = ["Общий фон", "Зима", "Осень", "Весна", "Лето"]
SEASONS = {"зима", "осень", "весна", "лето", "общий фон"}

def txt(c):
    return re.sub(r"\s+", " ", str(c.value)).strip() if c.value is not None else ""

def link(c):
    return c.hyperlink.target.strip() if c.hyperlink and c.hyperlink.target else ""

def authors_clean(s):
    s = re.sub(r"\[[^\]]*\]", "", s)      # [заказ, оплатил …]
    s = s.replace("(+)", "")
    s = re.sub(r"\s*[|+]\s*", ", ", s)
    return re.sub(r"\s+", " ", s).strip(" ,")

def variants_from(row, cols, headers, warn, name):
    out = []
    for col in cols:
        t, h = txt(row[col - 1]).lower(), link(row[col - 1])
        if h and (t in SEASONS):
            out.append({"label": t.capitalize(), "url": h})
        elif h:
            warn.append(f"{name}: ссылка без подписи сезона в колонке {headers[col]} -> подпись '{headers[col]}'")
            out.append({"label": headers[col], "url": h})
        elif t in SEASONS:
            warn.append(f"{name}: указан сезон «{t}», но в ячейке нет ссылки на файл")
    out.sort(key=lambda v: ORDER.index(v["label"]) if v["label"] in ORDER else 9)
    return out

wb = load_workbook(SRC)
items, missing, warn = [], [], []

# --- Актуальные ---
ws = wb["Фоны Актуальные"]
H1 = {4: "Зима", 5: "Осень", 6: "Весна", 7: "Лето"}
group = ""
for r, row in enumerate(ws.iter_rows(min_row=4), start=4):
    name = txt(row[1])
    if not name:
        continue
    if all(not txt(c) for i, c in enumerate(row) if i != 1):
        group = name.capitalize(); continue
    loc, status = txt(row[2]), txt(row[8])
    v = variants_from(row, [4, 5, 6, 7], H1, warn, name)
    base = {"name": name, "section": "Актуальные", "group": group, "location": loc}
    if v:
        items.append({**base, "authors": authors_clean(txt(row[7])), "variants": v})
    else:
        missing.append({**base, "reason": status or "нет ссылки на файл", "row": r})

# --- Архив старых ---
ws = wb["Фоны Архив старых"]
H2 = {3: "Зима", 4: "Осень", 5: "Весна", 6: "Лето"}
seen = OrderedDict()
for r, row in enumerate(ws.iter_rows(min_row=4), start=4):
    name = txt(row[1])
    if not name:
        continue
    v = variants_from(row, [3, 4, 5, 6], H2, warn, name + " (архив)")
    if not v:
        missing.append({"name": name, "section": "Архив старых", "group": "", "location": "", "reason": "нет ссылки на файл", "row": r}); continue
    key = (name, tuple(x["url"] for x in v))
    a = authors_clean(txt(row[6]))
    if key in seen:   # дубль строки с тем же файлом -> склеиваем авторов
        old = seen[key]["authors"]
        seen[key]["authors"] = ", ".join(x for x in dict.fromkeys([*old.split(", "), *a.split(", ")]) if x)
        warn.append(f"{name} (архив): строка {r} ссылается на тот же файл, что и предыдущая — объединено")
    else:
        seen[key] = {"name": name, "section": "Архив старых", "group": "", "location": "", "authors": a, "variants": v}
items += list(seen.values())

for i, it in enumerate(items, 1):
    it["id"] = i
data = {"built": datetime.date.today().isoformat(), "items": items, "missing": missing}
with open(OUT, "w", encoding="utf-8") as f:
    f.write("window.CATALOG = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n")

# --- отчёт ---
files = [v["url"] for it in items for v in it["variants"]]
print(f"Фонов в каталоге: {len(items)}; файлов-вариантов: {len(files)}; без файла: {len(missing)}")
exts = {}
for u in files:
    e = unquote(urlparse(u).path).rsplit(".", 1)[-1].lower(); exts[e] = exts.get(e, 0) + 1
print("Расширения:", exts)
print("Повторяющихся URL:", len(files) - len(set(files)))
odd = [u for u in set(files) if "/Design_Archive/refs/heads/main/" not in u]
print("Ссылки не вида Design_Archive/refs/heads/main:", len(odd))
for u in sorted(odd): print("  ", unquote(u))
print("\nБез файла:")
for m in missing: print(f"  [{m['section']}] {m['name']} ({m['location'] or '—'}) — {m['reason']}")
print("\nПредупреждения:")
for w in warn: print("  ", w)
